#!/usr/bin/env python3
"""Export Huaxi QIG quantization metrics to an Excel workbook."""

from __future__ import annotations

import argparse
import csv
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


DEFAULT_RUN_ROOT = Path("/work/model/lkp/experiments/qig/huaxi_qwen25vl_checkpoint1660_full_224_128")

METHOD_ORDER = [
    "fp16",
    "w4a16_rtn",
    "w4a16_gptq",
    "w4a16_awq",
    "w4a16_mbq",
    "w4a16_qig",
    "w4a8_rtn",
    "w4a8_smoothquant",
    "w4a8_smoothquant_alpha0p25_20260714_100302",
    "w4a8_smoothquant_alpha0p75_20260714_100302",
    "w4a8_mbq",
    "w4a8_qig",
    "w3a16_rtn",
    "w3a16_gptq",
    "w3a16_awq",
    "w3a16_mbq",
    "w3a16_qig",
    "w3a8_rtn",
    "w3a8_smoothquant",
    "w3a8_mbq",
    "w3a8_qig",
    "w2a16_awq",
    "w2a16_mbq",
    "w2a16_qig",
]

METRIC_RE = re.compile(
    r"^(.+?) - Average Precision: ([0-9.]+), Average Recall: ([0-9.]+), Average F1 Score: ([0-9.]+)"
)
CASE_RE = re.compile(
    r"^conclu_level1_case - Miss rate: ([0-9.]+), Case Recall@all: ([0-9.]+), Exact Match Accuracy: ([0-9.]+)"
)
DISEASE_RE = re.compile(
    r"^(conclu_level1_disease_(?:micro|macro)) - Sensitivity: ([0-9.]+), Specificity: ([0-9.]+), F1 Score: ([0-9.]+)"
)
GT_RE = re.compile(
    r"^gt_count=(.+?) \(n=([0-9]+)\) - Miss rate: ([0-9.]+), Case Recall@all: ([0-9.]+), Exact Match Accuracy: ([0-9.]+)"
)
LABEL_RE = re.compile(
    r"^(.+?) - Sensitivity: ([0-9.]+), Specificity: ([0-9.]+), F1 Score: ([0-9.]+), "
    r"tp: ([0-9]+), tn: ([0-9]+), fp: ([0-9]+), fn: ([0-9]+)"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-root", type=Path, default=DEFAULT_RUN_ROOT)
    parser.add_argument("--eval-root", type=Path, default=None)
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--csv-out", type=Path, default=None)
    parser.add_argument("--md-out", type=Path, default=None)
    return parser.parse_args()


def as_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def count_json_records(path: Path) -> int | None:
    if not path.exists():
        return None
    try:
        data = json.load(path.open("r", encoding="utf-8"))
    except Exception:
        return None
    if isinstance(data, dict) and isinstance(data.get("results"), list):
        return len(data["results"])
    if isinstance(data, list):
        return len(data)
    return None


def parse_method(name: str) -> dict[str, Any]:
    if name == "fp16":
        return {"family": "fp16", "method": "fp16", "w_bit": None, "a_bit": None, "variant": "baseline"}
    match = re.match(r"^w(\d+)a(\d+)_(.+)$", name)
    if not match:
        return {"family": "other", "method": name, "w_bit": None, "a_bit": None, "variant": ""}
    w_bit = int(match.group(1))
    a_bit = int(match.group(2))
    tail = match.group(3)
    if tail.startswith("smoothquant"):
        method = "smoothquant"
        variant = tail[len("smoothquant") :].lstrip("_")
    elif tail.startswith("rtn"):
        method = "rtn"
        variant = tail[len("rtn") :].lstrip("_")
    else:
        method = tail.split("_", 1)[0]
        variant = tail[len(method) :].lstrip("_")
    return {"family": f"w{w_bit}a{a_bit}", "method": method, "w_bit": w_bit, "a_bit": a_bit, "variant": variant}


def parse_metrics(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]]]:
    metrics: dict[str, Any] = {}
    by_gt: list[dict[str, Any]] = []
    per_label: list[dict[str, Any]] = []
    in_per_label = False
    if not path.exists():
        return metrics, by_gt, per_label

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("sample_count:"):
            metrics["sample_count"] = int(line.split(":", 1)[1].strip())
            continue
        if line == "conclu_level1_disease_per_label:":
            in_per_label = True
            continue
        match = METRIC_RE.match(line)
        if match:
            key = match.group(1)
            metrics[f"{key}_precision"] = as_float(match.group(2))
            metrics[f"{key}_recall"] = as_float(match.group(3))
            metrics[f"{key}_f1"] = as_float(match.group(4))
            continue
        match = DISEASE_RE.match(line)
        if match:
            key = match.group(1)
            metrics[f"{key}_sensitivity"] = as_float(match.group(2))
            metrics[f"{key}_specificity"] = as_float(match.group(3))
            metrics[f"{key}_f1"] = as_float(match.group(4))
            continue
        match = CASE_RE.match(line)
        if match:
            metrics["conclu_level1_case_miss_rate"] = as_float(match.group(1))
            metrics["conclu_level1_case_recall_all"] = as_float(match.group(2))
            metrics["conclu_level1_case_exact_match"] = as_float(match.group(3))
            continue
        match = GT_RE.match(line)
        if match:
            by_gt.append(
                {
                    "gt_count": match.group(1),
                    "n": int(match.group(2)),
                    "miss_rate": as_float(match.group(3)),
                    "case_recall_all": as_float(match.group(4)),
                    "exact_match": as_float(match.group(5)),
                }
            )
            continue
        if in_per_label:
            match = LABEL_RE.match(line)
            if match:
                per_label.append(
                    {
                        "label": match.group(1),
                        "sensitivity": as_float(match.group(2)),
                        "specificity": as_float(match.group(3)),
                        "f1": as_float(match.group(4)),
                        "tp": int(match.group(5)),
                        "tn": int(match.group(6)),
                        "fp": int(match.group(7)),
                        "fn": int(match.group(8)),
                    }
                )
    return metrics, by_gt, per_label


def file_info(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"exists": False, "size_mb": None, "mtime": None, "path": str(path)}
    stat = path.stat()
    return {
        "exists": True,
        "size_mb": round(stat.st_size / (1024 * 1024), 3),
        "mtime": datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "path": str(path),
    }


def original_status(out_dir: Path) -> str:
    done = (out_dir / "eval.done").exists()
    failed = (out_dir / "eval.failed").exists()
    if done and failed:
        return "DONE+FAILED"
    if done:
        return "DONE"
    if failed:
        return "FAILED"
    return "NONE"


def order_key(row: dict[str, Any]) -> tuple[int, Any]:
    name = row["method_id"]
    if name in METHOD_ORDER:
        return (0, METHOD_ORDER.index(name))
    if row.get("metrics_exists"):
        return (1, name)
    return (2, name)


def collect(eval_root: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    rows: list[dict[str, Any]] = []
    gt_rows: list[dict[str, Any]] = []
    label_rows: list[dict[str, Any]] = []

    for out_dir in sorted([path for path in eval_root.iterdir() if path.is_dir()]):
        name = out_dir.name
        raw = out_dir / "predictions_raw.json"
        pred = out_dir / "predictions.json"
        metrics_path = out_dir / "metrics.txt"
        metrics, by_gt, per_label = parse_metrics(metrics_path)
        parsed = parse_method(name)
        raw_info = file_info(raw)
        pred_info = file_info(pred)
        metrics_info = file_info(metrics_path)
        status_file = original_status(out_dir)

        if metrics_path.exists():
            usable_status = "METRICS_OK"
            if status_file == "FAILED":
                usable_status = "METRICS_OK_RECOMPUTED"
        elif raw.exists():
            usable_status = "RAW_ONLY"
        else:
            usable_status = "NO_RESULT"

        row: dict[str, Any] = {
            "method_id": name,
            "primary_included": "Y" if name in METHOD_ORDER else "N",
            "usable_status": usable_status,
            "original_status_file": status_file,
            "sample_count": metrics.get("sample_count") or count_json_records(raw),
            "raw_records": count_json_records(raw),
            "raw_exists": raw_info["exists"],
            "raw_size_mb": raw_info["size_mb"],
            "raw_mtime": raw_info["mtime"],
            "pred_exists": pred_info["exists"],
            "pred_size_mb": pred_info["size_mb"],
            "metrics_exists": metrics_info["exists"],
            "metrics_mtime": metrics_info["mtime"],
            "raw_path": str(raw),
            "pred_path": str(pred),
            "metrics_path": str(metrics_path),
        }
        row.update(parsed)
        row.update(metrics)
        rows.append(row)

        for item in by_gt:
            gt_rows.append({"method_id": name, "primary_included": row["primary_included"], **item})
        for item in per_label:
            label_rows.append(
                {
                    "method_id": name,
                    "primary_included": row["primary_included"],
                    "family": row["family"],
                    "method": row["method"],
                    "variant": row["variant"],
                    **item,
                }
            )

    fp16 = next((row for row in rows if row["method_id"] == "fp16"), None)
    if fp16:
        base_f1 = fp16.get("overall_level1_f1")
        base_case = fp16.get("conclu_level1_case_recall_all")
        for row in rows:
            f1 = row.get("overall_level1_f1")
            case = row.get("conclu_level1_case_recall_all")
            row["delta_overall_l1_f1_vs_fp16"] = f1 - base_f1 if f1 is not None and base_f1 else None
            row["overall_l1_f1_retention_vs_fp16"] = f1 / base_f1 if f1 is not None and base_f1 else None
            row["delta_case_recall_vs_fp16"] = case - base_case if case is not None and base_case else None

    return sorted(rows, key=order_key), gt_rows, label_rows


def write_csv_summary(
    rows: list[dict[str, Any]],
    path: Path,
    method_order: list[str] | None = None,
) -> None:
    metric_groups = [
        "overall_level1",
        "desc_level1",
        "conclu_level1",
        "overall_level2",
        "overall_level3",
    ]
    keys = ["method", "status", "sample_count"]
    for metric in metric_groups:
        keys.extend([f"{metric}_precision", f"{metric}_recall", f"{metric}_f1"])
    keys.extend(
        [
            "conclu_level1_disease_micro_sensitivity",
            "conclu_level1_disease_micro_specificity",
            "conclu_level1_disease_micro_f1",
            "conclu_level1_disease_macro_sensitivity",
            "conclu_level1_disease_macro_specificity",
            "conclu_level1_disease_macro_f1",
            "conclu_level1_case_miss_rate",
            "conclu_level1_case_recall_all",
            "conclu_level1_case_exact_match",
            "delta_overall_level1_f1_vs_fp16",
            "overall_level1_f1_retention_pct",
        ]
    )

    selected_methods = method_order or METHOD_ORDER
    primary_rows = [row for row in rows if row["method_id"] in selected_methods]
    primary_rows.sort(key=lambda row: selected_methods.index(row["method_id"]))
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        for row in primary_rows:
            out = {
                "method": row["method_id"],
                "status": "DONE" if row.get("metrics_exists") else row["usable_status"],
                "sample_count": row.get("sample_count"),
            }
            for key in keys[3:-2]:
                value = row.get(key)
                out[key] = f"{value:.4f}" if isinstance(value, float) else value
            delta = row.get("delta_overall_l1_f1_vs_fp16")
            retention = row.get("overall_l1_f1_retention_vs_fp16")
            out["delta_overall_level1_f1_vs_fp16"] = (
                f"{delta:.4f}" if delta is not None else ""
            )
            out["overall_level1_f1_retention_pct"] = (
                f"{retention * 100:.2f}" if retention is not None else ""
            )
            writer.writerow(out)


def write_md_summary(
    rows: list[dict[str, Any]],
    path: Path,
    method_order: list[str] | None = None,
) -> None:
    selected_methods = method_order or METHOD_ORDER
    primary_rows = [row for row in rows if row["method_id"] in selected_methods]
    primary_rows.sort(key=lambda row: selected_methods.index(row["method_id"]))
    headers = [
        "Method",
        "Status",
        "N",
        "Overall L1 F1",
        "Delta vs FP16",
        "Retention",
        "Desc L1 F1",
        "Conclu L1 F1",
        "Disease Micro F1",
        "Case Recall@all",
        "Case Exact",
    ]

    def fmt(value: Any) -> str:
        return f"{value:.4f}" if isinstance(value, float) else str(value or "")

    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        handle.write("| " + " | ".join(headers) + " |\n")
        handle.write("|" + "|".join(["---"] * len(headers)) + "|\n")
        for row in primary_rows:
            delta = row.get("delta_overall_l1_f1_vs_fp16")
            retention = row.get("overall_l1_f1_retention_vs_fp16")
            values = [
                row["method_id"],
                "DONE" if row.get("metrics_exists") else row["usable_status"],
                row.get("sample_count"),
                row.get("overall_level1_f1"),
                f"{delta:+.4f}" if delta is not None else "",
                f"{retention * 100:.2f}%" if retention is not None else "",
                row.get("desc_level1_f1"),
                row.get("conclu_level1_f1"),
                row.get("conclu_level1_disease_micro_f1"),
                row.get("conclu_level1_case_recall_all"),
                row.get("conclu_level1_case_exact_match"),
            ]
            handle.write("| " + " | ".join(fmt(value) for value in values) + " |\n")


def write_table(ws, rows: list[dict[str, Any]], cols: list[tuple[str, str]], table_name: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([title for _, title in cols])
    for row in rows:
        ws.append([row.get(key) for key, _ in cols])

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top")
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if table_name and ws.max_row >= 2:
        ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    for idx, (_, title) in enumerate(cols, start=1):
        width = max(len(title) + 2, 12)
        for row_idx in range(2, min(ws.max_row, 40) + 1):
            value = ws.cell(row=row_idx, column=idx).value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, 60))
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_workbook(rows: list[dict[str, Any]], gt_rows: list[dict[str, Any]], label_rows: list[dict[str, Any]], eval_root: Path) -> Workbook:
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme.append(["generated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    readme.append(["eval_root", str(eval_root)])
    readme.append(["note", "Rows with METRICS_OK_RECOMPUTED had metrics regenerated from existing raw predictions."])
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 120

    summary_cols = [
        ("method_id", "method_id"),
        ("usable_status", "usable_status"),
        ("original_status_file", "original_status_file"),
        ("family", "family"),
        ("method", "method"),
        ("variant", "variant"),
        ("w_bit", "w_bit"),
        ("a_bit", "a_bit"),
        ("sample_count", "sample_count"),
        ("overall_level1_precision", "overall_l1_p"),
        ("overall_level1_recall", "overall_l1_r"),
        ("overall_level1_f1", "overall_l1_f1"),
        ("desc_level1_f1", "desc_l1_f1"),
        ("conclu_level1_f1", "conclu_l1_f1"),
        ("overall_level2_f1", "overall_l2_f1"),
        ("overall_level3_f1", "overall_l3_f1"),
        ("conclu_level1_disease_micro_f1", "disease_micro_f1"),
        ("conclu_level1_disease_macro_f1", "disease_macro_f1"),
        ("conclu_level1_case_miss_rate", "case_miss_rate"),
        ("conclu_level1_case_recall_all", "case_recall_all"),
        ("conclu_level1_case_exact_match", "case_exact_match"),
        ("delta_overall_l1_f1_vs_fp16", "delta_overall_l1_f1_vs_fp16"),
        ("overall_l1_f1_retention_vs_fp16", "overall_l1_retention_vs_fp16"),
        ("delta_case_recall_vs_fp16", "delta_case_recall_vs_fp16"),
    ]
    all_cols = summary_cols + [
        ("primary_included", "primary_included"),
        ("raw_records", "raw_records"),
        ("raw_exists", "raw_exists"),
        ("raw_size_mb", "raw_size_mb"),
        ("raw_mtime", "raw_mtime"),
        ("pred_exists", "pred_exists"),
        ("pred_size_mb", "pred_size_mb"),
        ("metrics_exists", "metrics_exists"),
        ("metrics_mtime", "metrics_mtime"),
    ]

    primary_rows = [row for row in rows if row["method_id"] in METHOD_ORDER]
    ranking_rows = [dict(row, rank=rank) for rank, row in enumerate(
        sorted([row for row in primary_rows if row.get("overall_level1_f1") is not None], key=lambda item: item["overall_level1_f1"], reverse=True),
        start=1,
    )]
    best_rows = []
    for family in ["w4a16", "w4a8", "w3a16", "w3a8", "w2a16"]:
        candidates = [row for row in primary_rows if row["family"] == family and row.get("overall_level1_f1") is not None]
        if candidates:
            best_rows.append(max(candidates, key=lambda row: row["overall_level1_f1"]))

    ws = wb.create_sheet("Summary")
    write_table(ws, primary_rows, summary_cols, "SummaryTable")
    for title in ["overall_l1_f1", "conclu_l1_f1", "case_recall_all"]:
        col_idx = [header for _, header in summary_cols].index(title) + 1
        col = get_column_letter(col_idx)
        ws.conditional_formatting.add(
            f"{col}2:{col}{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
        )

    ws = wb.create_sheet("Ranking")
    write_table(ws, ranking_rows, [("rank", "rank")] + summary_cols, "RankingTable")

    ws = wb.create_sheet("Best_By_Family")
    write_table(ws, best_rows, summary_cols, "BestByFamilyTable")

    ws = wb.create_sheet("All_Runs")
    write_table(ws, rows, all_cols, "AllRunsTable")

    ws = wb.create_sheet("Paths")
    path_cols = [("method_id", "method_id"), ("usable_status", "usable_status"), ("raw_path", "raw_path"), ("pred_path", "pred_path"), ("metrics_path", "metrics_path")]
    write_table(ws, rows, path_cols, "PathsTable")
    for col in ["C", "D", "E"]:
        ws.column_dimensions[col].width = 105

    ws = wb.create_sheet("Case_By_GT_Count")
    write_table(
        ws,
        gt_rows,
        [("method_id", "method_id"), ("primary_included", "primary_included"), ("gt_count", "gt_count"), ("n", "n"), ("miss_rate", "miss_rate"), ("case_recall_all", "case_recall_all"), ("exact_match", "exact_match")],
        "CaseByGTTable",
    )

    ws = wb.create_sheet("Disease_Per_Label")
    write_table(
        ws,
        label_rows,
        [("method_id", "method_id"), ("primary_included", "primary_included"), ("family", "family"), ("method", "method"), ("variant", "variant"), ("label", "label"), ("sensitivity", "sensitivity"), ("specificity", "specificity"), ("f1", "f1"), ("tp", "tp"), ("tn", "tn"), ("fp", "fp"), ("fn", "fn")],
        "DiseasePerLabelTable",
    )
    return wb


def main() -> None:
    args = parse_args()
    eval_root = args.eval_root or args.run_root / "eval_test3100"
    out = args.out or eval_root / "quantization_results_summary_latest.xlsx"
    csv_out = args.csv_out or eval_root / "metrics_summary.csv"
    md_out = args.md_out or eval_root / "metrics_summary.md"
    rows, gt_rows, label_rows = collect(eval_root)
    wb = build_workbook(rows, gt_rows, label_rows, eval_root)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    write_csv_summary(rows, csv_out)
    write_md_summary(rows, md_out)
    w3w2_order = [
        "w3a16_rtn",
        "w3a16_gptq",
        "w3a16_awq",
        "w3a16_mbq",
        "w3a16_qig",
        "w2a16_awq",
        "w2a16_mbq",
        "w2a16_qig",
    ]
    w3w2_csv = eval_root / "metrics_summary_w3w2_a16.csv"
    w3w2_md = eval_root / "metrics_summary_w3w2_a16.md"
    write_csv_summary(rows, w3w2_csv, w3w2_order)
    write_md_summary(rows, w3w2_md, w3w2_order)
    print(f"[OK] wrote {out}")
    print(f"[OK] wrote {csv_out}")
    print(f"[OK] wrote {md_out}")
    print(f"[OK] wrote {w3w2_csv}")
    print(f"[OK] wrote {w3w2_md}")
    print(f"[OK] rows={len(rows)} gt_rows={len(gt_rows)} label_rows={len(label_rows)}")


if __name__ == "__main__":
    main()
